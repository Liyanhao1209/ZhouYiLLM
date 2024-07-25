from text_split.QA_split.splitter.rules.abRule import abRule
from docx import Document
import logging


class word_limit_Rule(abRule):

    def __init__(self, word_limit: int):
        self.wl = word_limit

    def split(self, file_path: str) -> list:
        doc = Document(file_path)
        pages = []
        curr_page = []
        size = 0

        for para in doc.paragraphs:
            if size + len(para.text) <= self.wl:
                curr_page.append(para.text)
                size += len(para.text)
            else:
                pages.append(curr_page)
                curr_page = [para.text]
                size = 0

        if curr_page:
            pages.append(curr_page)

        return pages

    def save(self, target_path: str, articles: list):
        for i, article in enumerate(articles):
            new_doc = Document()
            for para in article:
                new_doc.add_paragraph(para)
            new_doc.save(f"{target_path}/article_{i}.docx")
